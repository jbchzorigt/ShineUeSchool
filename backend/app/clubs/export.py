"""Ээлжийн бүртгэлийг Excel (.xlsx) болгох: дугуйлан бүрт нэг sheet, мөрүүд ангиар (1→12) эрэмбэлэгдэнэ."""

import io
import re

from openpyxl import Workbook
from openpyxl.styles import Font

from .models import REG_CONFIRMED, ClubRegistration

HEADERS = ("Анги", "Сурагчийн овог", "Сурагчийн нэр", "Бүртгүүлэгчийн овог", "Бүртгүүлэгчийн нэр",
           "Утас", "Имэйл", "Төлөв", "Төлбөр", "Огноо")
_BAD = re.compile(r"[\\/*?:\[\]]")   # Excel-ийн sheet нэрэнд хориотой тэмдэгтүүд (\ / * ? : [ ])


def sheet_title(name: str, used: set[str]) -> str:
    """Дугуйлангийн нэрийг sheet-ийн нэр болгоно: хориотой тэмдэгт хасаж, 31 тэмдэгтээр таслаж, давхардвал (2), (3)…"""
    base = (_BAD.sub("", name).strip() or "Дугуйлан")[:31]
    title, n = base, 2
    while title in used:
        suffix = f" ({n})"
        title = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(title)
    return title


def build_registrations_xlsx(clubs: list[tuple[str, bool, list[ClubRegistration]]]) -> bytes:
    """clubs: (дугуйлангийн нэр, төлбөртэй эсэх, бүртгэлүүд) — ээлжийн дарааллаар; бүртгэлгүй бол толгойтой хоосон sheet."""
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for name, is_paid, regs in clubs:
        ws = wb.create_sheet(sheet_title(name, used))
        ws.append(HEADERS)
        for c in ws[1]:
            c.font = Font(bold=True)
        # Анги 1→12, дараа нь бүртгэгдсэн (хасагдсаны өмнө), дараа нь огноо
        for r in sorted(regs, key=lambda r: (r.grade, r.status != REG_CONFIRMED, r.created_at or 0, r.id)):
            paid = "—" if not is_paid else ("Төлсөн" if r.is_paid_marked else "Төлөөгүй")
            ws.append((r.grade, r.student_last_name, r.student_first_name, r.guardian_last_name, r.guardian_first_name,
                       r.phone, r.email, "Бүртгэгдсэн" if r.status == REG_CONFIRMED else "Хасагдсан", paid,
                       r.created_at.strftime("%Y-%m-%d %H:%M") if r.created_at else ""))
        for col, width in zip("ABCDEFGHIJ", (6, 16, 16, 18, 18, 14, 28, 12, 10, 17)):
            ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
