"""Дугуйлан — менежерийн API."""

from datetime import UTC, datetime, timedelta

from tests.helpers import manager_headers, seed_clubs, staff_headers

A = "/api/clubs/admin"


def club_payload(**over):
    now = datetime.now(UTC)
    d = {"name": "Хөл бөмбөг", "description": "Долоо хоногт 2 удаа",
         "quotas": [{"grade": 7, "capacity": 10}, {"grade": 5, "capacity": 5}],
         "is_paid": False, "fee": 0, "fee_note": "",
         "registration_start": (now - timedelta(days=1)).isoformat(),
         "registration_end": (now + timedelta(days=10)).isoformat(), "is_published": True}
    d.update(over)
    return d


async def test_admin_requires_manager(client, db, make_user):
    assert (await client.get(f"{A}/rounds/")).status_code == 401
    h = await staff_headers(client, make_user, "oly", roles=("olympiad",))
    assert (await client.get(f"{A}/rounds/", headers=h)).status_code == 403
    assert (await client.post(f"{A}/rounds/", headers=h, json={"name": "x"})).status_code == 403
    assert (await client.get(f"{A}/rounds/1/registrations.xlsx", headers=h)).status_code == 403
    h = await manager_headers(client, make_user)
    assert (await client.get(f"{A}/rounds/", headers=h)).status_code == 200


async def test_rounds_crud_and_single_active(client, db, make_user):
    h = await manager_headers(client, make_user)
    r = await client.post(f"{A}/rounds/", headers=h, json={"name": "  2026 намар "})
    assert r.status_code == 201 and r.json()["name"] == "2026 намар" and r.json()["is_active"] is False
    r1 = r.json()["id"]
    r2 = (await client.post(f"{A}/rounds/", headers=h, json={"name": "2027 хавар"})).json()["id"]
    assert (await client.post(f"{A}/rounds/", headers=h, json={"name": " "})).status_code == 400
    assert (await client.patch(f"{A}/rounds/{r1}/", headers=h, json={"is_active": True})).json()["is_active"] is True
    r = await client.patch(f"{A}/rounds/{r2}/", headers=h, json={"is_active": True, "name": "2027 хавар (шинэ)"})
    assert r.json()["is_active"] is True and r.json()["name"] == "2027 хавар (шинэ)"
    rounds = {x["id"]: x for x in (await client.get(f"{A}/rounds/", headers=h)).json()}
    assert rounds[r1]["is_active"] is False and rounds[r2]["is_active"] is True
    assert rounds[r2]["clubs_count"] == 0 and rounds[r2]["registrations_count"] == 0
    assert (await client.delete(f"{A}/rounds/{r1}/", headers=h)).status_code == 204
    assert (await client.patch(f"{A}/rounds/999999/", headers=h, json={"name": "x"})).status_code == 404


async def test_round_delete_rules(client, db, make_user):
    from app.clubs.models import ClubRegistration
    s = await seed_clubs(db)
    h = await manager_headers(client, make_user)
    rounds = (await client.get(f"{A}/rounds/", headers=h)).json()
    assert rounds[0]["clubs_count"] == 5
    db.add(ClubRegistration(club_id=s.chess_id, round_id=s.round_id, email="a@shineue.edu.mn",
                            student_last_name="Б", student_first_name="Бат", guardian_last_name="Д",
                            guardian_first_name="Дорж", phone="99001122", grade=5))
    await db.flush()
    r = await client.delete(f"{A}/rounds/{s.round_id}/", headers=h)
    assert r.status_code == 409 and "Бүртгэлтэй" in r.json()["detail"]
    r = await client.delete(f"{A}/clubs/{s.chess_id}/", headers=h)
    assert r.status_code == 409
    assert (await client.delete(f"{A}/clubs/{s.robot_id}/", headers=h)).status_code == 204


async def test_round_delete_with_clubs_no_registrations(client, db, make_user, tmp_path, monkeypatch):
    from app.config import settings
    monkeypatch.setattr(settings, "media_dir", tmp_path)
    s = await seed_clubs(db)
    h = await manager_headers(client, make_user)
    assert (await client.delete(f"{A}/rounds/{s.round_id}/", headers=h)).status_code == 204
    assert (await client.get(f"{A}/rounds/", headers=h)).json() == []
    assert (await client.get("/api/clubs/")).json() == {"round": None, "clubs": []}


async def test_clubs_create_validate_patch(client, db, make_user):
    s = await seed_clubs(db)
    h = await manager_headers(client, make_user)
    r = await client.post(f"{A}/rounds/{s.round_id}/clubs/", headers=h, json=club_payload())
    assert r.status_code == 201, r.text
    c = r.json()
    assert c["grades"] == [5, 7] and c["capacity"] == 15 and c["order"] == 6 and c["state"] == "open" and c["taken"] == 0
    assert c["quotas"] == [{"grade": 5, "capacity": 5, "taken": 0, "slots_left": 5, "full": False},
                           {"grade": 7, "capacity": 10, "taken": 0, "slots_left": 10, "full": False}]
    assert c["is_published"] and c["round_id"] == s.round_id and c["fee"] == 0
    cid = c["id"]
    # шалгалтууд
    now = datetime.now(UTC)
    r = await client.post(f"{A}/rounds/{s.round_id}/clubs/", headers=h,
                          json=club_payload(registration_start=(now + timedelta(days=2)).isoformat(),
                                            registration_end=(now + timedelta(days=1)).isoformat()))
    assert r.status_code == 400 and "registration_end" in r.json()
    r = await client.post(f"{A}/rounds/{s.round_id}/clubs/", headers=h, json=club_payload(is_paid=True, fee=0))
    assert r.status_code == 400 and "fee" in r.json()
    r = await client.post(f"{A}/rounds/{s.round_id}/clubs/", headers=h, json=club_payload(quotas=[]))
    assert r.status_code == 400 and "quotas" in r.json()
    r = await client.post(f"{A}/rounds/{s.round_id}/clubs/", headers=h, json=club_payload(quotas=[{"grade": 13, "capacity": 1}]))
    assert r.status_code == 400 and "quotas" in r.json()
    r = await client.post(f"{A}/rounds/{s.round_id}/clubs/", headers=h, json=club_payload(quotas=[{"grade": 5, "capacity": 0}]))
    assert r.status_code == 400 and "quotas" in r.json()
    r = await client.post(f"{A}/rounds/{s.round_id}/clubs/", headers=h,
                          json=club_payload(quotas=[{"grade": 5, "capacity": 1}, {"grade": 5, "capacity": 2}]))
    assert r.status_code == 400 and "quotas" in r.json()
    r = await client.post(f"{A}/rounds/{s.round_id}/clubs/", headers=h,
                          json=club_payload(quotas=[{"grade": 5, "capacity": 2.5}]))
    assert r.status_code == 400 and "quotas" in r.json()
    # төлбөргүй болгоход fee, fee_note цэвэрлэгдэнэ
    r = await client.post(f"{A}/rounds/{s.round_id}/clubs/", headers=h, json=club_payload(is_paid=False, fee=500, fee_note="сард"))
    assert r.json()["fee"] == 0 and r.json()["fee_note"] == ""
    # PATCH
    r = await client.patch(f"{A}/clubs/{cid}/", headers=h, json={"is_paid": True, "fee": 90000, "fee_note": "улиралд"})
    assert r.status_code == 200 and r.json()["fee"] == 90000
    r = await client.patch(f"{A}/clubs/{cid}/", headers=h, json={"fee": 0})            # is_paid хэвээр true → 400
    assert r.status_code == 400 and "fee" in r.json()
    r = await client.patch(f"{A}/clubs/{cid}/", headers=h, json={"registration_end": (now - timedelta(days=5)).isoformat()})
    assert r.status_code == 400 and "registration_end" in r.json()
    r = await client.patch(f"{A}/clubs/{cid}/", headers=h, json={"is_published": False, "quotas": [{"grade": 9, "capacity": 3}]})
    assert r.json()["is_published"] is False and r.json()["grades"] == [9] and r.json()["capacity"] == 3
    assert (await client.patch(f"{A}/clubs/999999/", headers=h, json={"name": "x"})).status_code == 404


async def test_club_quota_not_below_taken(client, db, make_user):
    from app.clubs.models import ClubRegistration
    s = await seed_clubs(db)
    h = await manager_headers(client, make_user)
    ids = []
    for i in range(2):
        reg = ClubRegistration(club_id=s.chess_id, round_id=s.round_id, email=f"s{i}@shineue.edu.mn",
                               student_last_name="Б", student_first_name="Бат", guardian_last_name="Д",
                               guardian_first_name="Дорж", phone="99001122", grade=5)
        db.add(reg)
        await db.flush()
        ids.append(reg.id)
    q = lambda five: [{"grade": 5, "capacity": five}, {"grade": 6, "capacity": 1}, {"grade": 7, "capacity": 1}, {"grade": 8, "capacity": 1}]
    r = await client.patch(f"{A}/clubs/{s.chess_id}/", headers=h, json={"quotas": q(1)})
    assert r.status_code == 400 and r.json()["quotas"] == ["5-р ангид бүртгэгдсэн 2 сурагчаас бага байж болохгүй"]
    r = await client.patch(f"{A}/clubs/{s.chess_id}/", headers=h, json={"quotas": [{"grade": 6, "capacity": 1}]})
    assert r.status_code == 400 and r.json()["quotas"] == ["5-р ангид бүртгэл байгаа тул хасах боломжгүй"]
    r = await client.patch(f"{A}/clubs/{s.chess_id}/", headers=h, json={"quotas": q(2)})
    assert r.status_code == 200 and r.json()["state"] == "open" and r.json()["quotas"][0]["full"] is True
    # нэгийг нь хассаны дараа квот 1 боломжтой
    assert (await client.post(f"{A}/registrations/{ids[0]}/remove/", headers=h)).status_code == 200
    r = await client.patch(f"{A}/clubs/{s.chess_id}/", headers=h, json={"quotas": q(1)})
    assert r.status_code == 200 and r.json()["taken"] == 1 and r.json()["quotas"][0]["full"] is True


async def test_clubs_list_and_order(client, db, make_user):
    s = await seed_clubs(db)
    h = await manager_headers(client, make_user)
    lst = (await client.get(f"{A}/rounds/{s.round_id}/clubs/", headers=h)).json()
    assert [c["name"] for c in lst] == ["Шатар", "Робот", "Дуу", "Зураг", "Нууц"]     # нийтлэгдээгүй ч орно
    ids = [c["id"] for c in lst]
    r = await client.put(f"{A}/rounds/{s.round_id}/clubs/order/", headers=h, json={"ids": list(reversed(ids))})
    assert r.status_code == 200 and [c["name"] for c in r.json()] == ["Нууц", "Зураг", "Дуу", "Робот", "Шатар"]
    r = await client.put(f"{A}/rounds/{s.round_id}/clubs/order/", headers=h, json={"ids": ids[:2]})
    assert r.status_code == 400 and "ids" in r.json()


import io

from PIL import Image


def png_bytes(color=(30, 58, 143)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (40, 30), color).save(buf, format="PNG")
    return buf.getvalue()


async def add_reg(db, s, club_id, email, grade, **over):
    from app.clubs.models import ClubRegistration
    d = dict(club_id=club_id, round_id=s.round_id, email=email, student_last_name="Бат", student_first_name="Дорж",
             guardian_last_name="Дорж", guardian_first_name="Сүх", phone="99001122", grade=grade)
    d.update(over)
    reg = ClubRegistration(**d)
    db.add(reg)
    await db.flush()
    return reg.id


async def test_club_images_upload_order_delete(client, db, make_user, tmp_path, monkeypatch):
    from app.config import settings
    monkeypatch.setattr(settings, "media_dir", tmp_path)
    s = await seed_clubs(db)
    h = await manager_headers(client, make_user)
    r = await client.post(f"{A}/clubs/{s.chess_id}/images/", headers=h, files={"image": ("a.png", png_bytes(), "image/png")})
    assert r.status_code == 201, r.text
    r = await client.post(f"{A}/clubs/{s.chess_id}/images/", headers=h, files={"image": ("b.png", png_bytes((255, 0, 0)), "image/png")})
    imgs = r.json()["images"]
    assert len(imgs) == 2 and [i["order"] for i in imgs] == [1, 2] and "/media/clubs/" in imgs[0]["url"]
    rel0 = imgs[0]["url"].split("/media/")[1]
    assert (tmp_path / rel0).exists()
    r = await client.post(f"{A}/clubs/{s.chess_id}/images/", headers=h, files={"image": ("t.txt", b"hello", "text/plain")})
    assert r.status_code == 400 and "image" in r.json()
    # эрэмбэ
    r = await client.put(f"{A}/clubs/{s.chess_id}/images/order/", headers=h, json={"ids": [imgs[1]["id"], imgs[0]["id"]]})
    assert r.status_code == 200 and [i["id"] for i in r.json()["images"]] == [imgs[1]["id"], imgs[0]["id"]]
    r = await client.put(f"{A}/clubs/{s.chess_id}/images/order/", headers=h, json={"ids": [imgs[0]["id"]]})
    assert r.status_code == 400
    # олон нийтэд харагдана
    pub = {c["name"]: c for c in (await client.get("/api/clubs/")).json()["clubs"]}
    assert len(pub["Шатар"]["images"]) == 2
    # устгах
    r = await client.delete(f"{A}/images/{imgs[0]['id']}/", headers=h)
    assert r.status_code == 200 and len(r.json()["images"]) == 1 and not (tmp_path / rel0).exists()
    assert (await client.delete(f"{A}/images/{imgs[0]['id']}/", headers=h)).status_code == 404
    # дугуйлан устгахад файл устана
    rel1 = r.json()["images"][0]["url"].split("/media/")[1]
    assert (await client.delete(f"{A}/clubs/{s.chess_id}/", headers=h)).status_code == 204
    assert not (tmp_path / rel1).exists()


async def test_registrations_list_remove_paid(client, db, make_user):
    s = await seed_clubs(db)
    h = await manager_headers(client, make_user)
    r1 = await add_reg(db, s, s.chess_id, "a@shineue.edu.mn", 5)
    r2 = await add_reg(db, s, s.chess_id, "b@shineue.edu.mn", 6, status="removed")
    r3 = await add_reg(db, s, s.robot_id, "c@shineue.edu.mn", 9)
    lst = (await client.get(f"{A}/clubs/{s.chess_id}/registrations/", headers=h)).json()
    assert [x["id"] for x in lst] == [r1, r2] and lst[0]["status"] == "confirmed" and lst[1]["status"] == "removed"
    assert lst[0]["email"] == "a@shineue.edu.mn" and lst[0]["is_paid_marked"] is False and lst[0]["removed_at"] is None
    # төлсөн тэмдэглэх
    r = await client.patch(f"{A}/registrations/{r3}/", headers=h, json={"is_paid_marked": True})
    assert r.status_code == 200 and r.json()["is_paid_marked"] is True
    # хасах → слот суларна
    r = await client.post(f"{A}/registrations/{r1}/remove/", headers=h)
    assert r.status_code == 200 and r.json()["status"] == "removed" and r.json()["removed_at"]
    assert (await client.post(f"{A}/registrations/{r1}/remove/", headers=h)).status_code == 409
    pub = {c["name"]: c for c in (await client.get("/api/clubs/")).json()["clubs"]}
    assert pub["Шатар"]["taken"] == 0
    assert (await client.get(f"{A}/clubs/999999/registrations/", headers=h)).status_code == 404


async def test_registrations_xlsx(client, db, make_user):
    from openpyxl import load_workbook
    s = await seed_clubs(db)
    h = await manager_headers(client, make_user)
    await add_reg(db, s, s.chess_id, "d@shineue.edu.mn", 8)                    # эхэлж нэмэгдсэн ч ангиар хойно
    await add_reg(db, s, s.chess_id, "a@shineue.edu.mn", 5)
    await add_reg(db, s, s.robot_id, "c@shineue.edu.mn", 9, is_paid_marked=True)
    await add_reg(db, s, s.chess_id, "b@shineue.edu.mn", 6, status="removed")
    r = await client.get(f"{A}/rounds/{s.round_id}/registrations.xlsx", headers=h)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "attachment" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Шатар", "Робот", "Дуу", "Зураг", "Нууц"]        # дугуйлан бүрт sheet, ээлжийн дарааллаар
    header = ("Анги", "Сурагчийн овог", "Сурагчийн нэр", "Бүртгүүлэгчийн овог", "Бүртгүүлэгчийн нэр",
              "Утас", "Имэйл", "Төлөв", "Төлбөр", "Огноо")
    rows = list(wb["Шатар"].iter_rows(values_only=True))
    assert rows[0] == header and len(rows) == 4
    assert [r[0] for r in rows[1:]] == [5, 6, 8]                                   # ангиар 1→12
    by_email = {r[6]: r for r in rows[1:]}
    assert by_email["a@shineue.edu.mn"][7] == "Бүртгэгдсэн" and by_email["a@shineue.edu.mn"][8] == "—"
    assert by_email["b@shineue.edu.mn"][7] == "Хасагдсан"
    robot = list(wb["Робот"].iter_rows(values_only=True))
    assert len(robot) == 2 and robot[1][8] == "Төлсөн" and robot[1][0] == 9
    assert list(wb["Дуу"].iter_rows(values_only=True)) == [header]               # бүртгэлгүй дугуйлан: зөвхөн толгой
    assert (await client.get(f"{A}/rounds/999999/registrations.xlsx", headers=h)).status_code == 404
